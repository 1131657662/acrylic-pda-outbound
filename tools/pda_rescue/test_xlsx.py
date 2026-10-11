"""Workbook compatibility uses synthetic files and checks final values, not authoring internals."""
import json
import os
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from .xlsx_text import repair_and_verify


class RescueWorkbookTest(unittest.TestCase):
    def test_sheet_relationships_and_dangerous_text_are_preserved(self):
        directory = Path(tempfile.mkdtemp(prefix="xlsx-test-", dir=os.environ["PDA_RESCUE_VECTOR_DIR"]))
        path = directory / "synthetic.xlsx"
        book = Workbook()
        book.active.title = "第一"
        book.create_sheet("第二")
        for sheet in book:
            sheet.append(["header"])
            sheet["A5"] = "coerced"
            sheet["B5"] = 7
        book.save(path)
        expected = [{"name": "第二", "rows": [["=danger\n003", 7]]},
                    {"name": "第一", "rows": [["9007199254740993", 7]]}]
        Path(str(path) + ".expected.json").write_text(json.dumps(expected), encoding="utf-8")
        repair_and_verify(path)
        result = load_workbook(path)
        self.assertEqual("9007199254740993", result["第一"]["A5"].value)
        self.assertEqual("=danger\n003", result["第二"]["A5"].value)
        self.assertEqual("s", result["第二"]["A5"].data_type)
        result.close()


if __name__ == "__main__":
    unittest.main()
