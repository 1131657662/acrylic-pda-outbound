"""Standard preview reads the final XLSX, so string IDs cannot be coerced by the authoring renderer."""
import os
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


def render(path: Path):
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        font_path = os.environ.get("PDA_RESCUE_FONT", "C:/Windows/Fonts/msyh.ttc")
        font = ImageFont.truetype(font_path, 17)
        for sheet in book:
            sheet.calculate_dimension(force=True)
            _sheet(path, sheet, font)
    finally:
        book.close()


def _sheet(path, sheet, font):
    rows = [list(row) for row in sheet.iter_rows(min_row=4, max_row=min(sheet.max_row, 20), values_only=True)]
    texts = [["" if value is None else str(value) for value in row] for row in rows]
    probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    widths = [min(600, max(130, max(int(probe.textlength(row[col][:160], font=font)) + 24 for row in texts))) for col in range(sheet.max_column)]
    image = Image.new("RGB", (sum(widths) + 40, len(texts) * 36 + 90), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 12), sheet.title, font=font, fill="#1F2937")
    draw.text((20, 36), "前16行；长文本截显示；完整值见Excel/原件", font=font, fill="#1F2937")
    for row_index, row in enumerate(texts):
        x, y = 20, 65 + row_index * 36
        for column, text in enumerate(row):
            fill = "#294862" if row_index == 0 else ("#F0F4F7" if row_index % 2 else "#FFFFFF")
            draw.rectangle((x, y, x + widths[column], y + 35), fill=fill)
            visible = _fit(text, widths[column] - 16, lambda value: probe.textlength(value, font=font))
            draw.text((x + 8, y + 6), visible, font=font, fill="white" if row_index == 0 else "#1F2937")
            x += widths[column]
    image.save(str(path) + f".{sheet.title}.verified.png")


def _fit(text, width, measure):
    if len(text) <= 160 and measure(text) <= width:
        return text
    value = text[:160]
    while value and measure(value + "…") > width:
        value = value[:-1]
    return value + "…"
